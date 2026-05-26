from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.escopo_categorias import CATEGORIA_IDENTIFICACAO
from src.validar_resultados import (
    COLUNAS_FONTES_LOG,
    COLUNAS_MUNICIPIOS,
    COLUNAS_PENDENCIAS,
    COLUNAS_RESULTADO,
    garantir_colunas,
)
from src.utils import load_yaml, project_path


STATUS_COM_DADO = {"Encontrado", "Parcial"}
CAMPOS_DADO_PUBLICADO = ["Nome", "E-mail", "Telefone", "Celular/WhatsApp", "Celular"]
OBS_SEM_DADO_PUBLICADO = "Páginas oficiais consultadas, mas sem dado oficial claro para esta categoria."
FONTES_INDISPONIVEIS_VALIDACAO = ("antonina.pr.gov.br/secretariaview",)
OBS_FONTE_INDISPONIVEL_VALIDACAO = (
    "Fonte oficial retornou indisponibilidade durante a validação; dado publicado previamente não foi mantido sem confirmação atual."
)


def _valor_preenchido(value: object) -> bool:
    if pd.isna(value):
        return False
    text = str(value).strip()
    return bool(text) and text.lower() not in {"nan", "none", "null"}


def _juntar_observacao(atual: object, nova: str) -> str:
    atual_texto = "" if pd.isna(atual) else str(atual).strip()
    if not atual_texto:
        return nova
    if nova in atual_texto:
        return atual_texto
    return f"{atual_texto} {nova}"


def _fonte_indisponivel_validacao(row: pd.Series) -> bool:
    url = str(row.get("URL da fonte", "") or row.get("URL específica", "") or "").lower()
    return any(token in url for token in FONTES_INDISPONIVEIS_VALIDACAO)


def normalizar_status_sem_dado_publicado(resultado_df: pd.DataFrame) -> pd.DataFrame:
    resultado = resultado_df.copy()
    if resultado.empty or "Status" not in resultado.columns:
        return resultado

    for column in CAMPOS_DADO_PUBLICADO + ["Cargo/Área", "Observações"]:
        if column not in resultado.columns:
            resultado[column] = ""

    try:
        from src.extrair_contatos import _nome_valido
    except Exception:  # pragma: no cover - proteção contra import parcial em empacotamentos
        _nome_valido = None

    for index, row in resultado.iterrows():
        cargo = str(row.get("Cargo/Área", "") or "").strip()
        status = str(row.get("Status", "") or "").strip()
        if cargo == CATEGORIA_IDENTIFICACAO.label or status not in STATUS_COM_DADO:
            continue

        if _fonte_indisponivel_validacao(row):
            for campo in CAMPOS_DADO_PUBLICADO:
                resultado.at[index, campo] = ""
            resultado.at[index, "Status"] = "Site fora do ar"
            resultado.at[index, "Observações"] = _juntar_observacao(
                row.get("Observações", ""),
                OBS_FONTE_INDISPONIVEL_VALIDACAO,
            )
            continue

        nome = row.get("Nome", "")
        if _valor_preenchido(nome) and _nome_valido is not None and not _nome_valido(str(nome)):
            resultado.at[index, "Nome"] = ""
            resultado.at[index, "Observações"] = _juntar_observacao(
                row.get("Observações", ""),
                "Nome descartado por parecer rótulo institucional, navegação ou texto sem pessoa publicada.",
            )

        dados_presentes = any(_valor_preenchido(resultado.at[index, campo]) for campo in CAMPOS_DADO_PUBLICADO)
        if dados_presentes:
            continue

        resultado.at[index, "Status"] = "Não publicado"
        resultado.at[index, "Observações"] = _juntar_observacao(
            resultado.at[index, "Observações"],
            OBS_SEM_DADO_PUBLICADO,
        )

    return resultado


def criar_resumo(
    resultado_df: pd.DataFrame,
    municipios_df: pd.DataFrame,
    pendencias_df: pd.DataFrame,
) -> pd.DataFrame:
    total_municipios = len(municipios_df)
    resultado_contatos = _resultado_sem_identificacao(resultado_df)
    encontrados = resultado_contatos[resultado_contatos.get("Status", "") == "Encontrado"] if not resultado_contatos.empty else resultado_contatos
    parciais = resultado_contatos[resultado_contatos.get("Status", "") == "Parcial"] if not resultado_contatos.empty else resultado_contatos

    status_geral = municipios_df.get("Status geral", pd.Series(dtype=str)).astype(str)
    pend_status = pendencias_df.get("Status", pd.Series(dtype=str)).astype(str)

    rows = [
        _resumo_row("Indicadores gerais", "Total de municípios no escopo", total_municipios),
        _resumo_row("Indicadores gerais", "Total de municípios com algum dado encontrado", _contar_municipios_unicos(encontrados)),
        _resumo_row("Indicadores gerais", "Total de municípios com dados parciais", _contar_municipios_unicos(parciais)),
        _resumo_row(
            "Indicadores gerais",
            "Total de municípios sem dados encontrados",
            int(status_geral.isin(["Não encontrado", "Não publicado", "Página sem informação pública"]).sum()),
        ),
        _resumo_row(
            "Indicadores gerais",
            "Total de sites não localizados",
            int((status_geral == "Site não localizado").sum() or (pend_status == "Site não localizado").sum()),
        ),
        _resumo_row(
            "Indicadores gerais",
            "Total de sites fora do ar",
            int((status_geral == "Site fora do ar").sum() or (pend_status == "Site fora do ar").sum()),
        ),
        _resumo_row(
            "Indicadores gerais",
            "Total de bloqueios técnicos",
            int((status_geral == "Bloqueio técnico").sum() or (pend_status == "Bloqueio técnico").sum()),
        ),
        _resumo_row("Indicadores gerais", "Data/hora da execução", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    rows.extend(_resumo_agrupado(municipios_df, "Por UF", ["UF"], "Status geral"))
    rows.extend(_resumo_agrupado(municipios_df, "Por Município/Capital", ["UF", "Município/Capital", "Esfera"], "Status geral"))
    rows.extend(_resumo_agrupado(municipios_df, "Por Esfera", ["Esfera"], "Status geral"))
    rows.extend(_resumo_agrupado(resultado_df, "Por Cargo/Área", ["Cargo/Área"], "Status"))
    rows.extend(_resumo_agrupado(resultado_df, "Por Status", [], "Status"))

    columns = ["Seção", "Indicador", "Valor", "UF", "Município/Capital", "Esfera", "Cargo/Área", "Status", "Quantidade"]
    return pd.DataFrame(rows, columns=columns)


def _resumo_row(
    secao: str,
    indicador: str = "",
    valor: object = "",
    uf: object = "",
    municipio_capital: object = "",
    esfera: object = "",
    cargo_area: object = "",
    status: object = "",
    quantidade: object = "",
) -> dict[str, object]:
    return {
        "Seção": secao,
        "Indicador": indicador,
        "Valor": valor,
        "UF": uf,
        "Município/Capital": municipio_capital,
        "Esfera": esfera,
        "Cargo/Área": cargo_area,
        "Status": status,
        "Quantidade": quantidade,
    }


def _resumo_agrupado(
    df: pd.DataFrame,
    secao: str,
    group_columns: list[str],
    status_column: str,
) -> list[dict[str, object]]:
    columns = list(group_columns) + [status_column]
    if df.empty:
        return [_resumo_row(secao, "Sem dados", 0)]

    work = df.copy()
    for column in columns:
        if column not in work.columns:
            work[column] = ""
    work[columns] = work[columns].fillna("").astype(str)
    grouped = work.groupby(columns, dropna=False).size().reset_index(name="Quantidade")

    rows: list[dict[str, object]] = []
    for _, row in grouped.iterrows():
        rows.append(
            _resumo_row(
                secao=secao,
                uf=row.get("UF", ""),
                municipio_capital=row.get("Município/Capital", ""),
                esfera=row.get("Esfera", ""),
                cargo_area=row.get("Cargo/Área", ""),
                status=row.get(status_column, ""),
                quantidade=int(row["Quantidade"]),
            )
        )
    return rows


def _contar_municipios_unicos(df: pd.DataFrame) -> int:
    if df.empty or not {"UF", "Município"}.issubset(df.columns):
        return 0
    return int(df[["UF", "Município"]].drop_duplicates().shape[0])


def _resultado_sem_identificacao(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or "Cargo/Área" not in df.columns:
        return df
    return df[df["Cargo/Área"].astype(str) != CATEGORIA_IDENTIFICACAO.label].copy()


def _formatar_aba(writer: pd.ExcelWriter, sheet_name: str) -> None:
    worksheet = writer.book[sheet_name]
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


def criar_configuracao_df() -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    for nome, relative in {
        "cargos": ("config", "cargos.yml"),
        "palavras_chave": ("config", "palavras_chave.yml"),
        "scraping": ("config", "scraping.yml"),
        "estados": ("config", "estados.yml"),
        "governos_estaduais": ("config", "governos_estaduais.yml"),
    }.items():
        path = project_path(*relative)
        if not path.exists():
            rows.append({"Arquivo": str(path), "Grupo": nome, "Chave": "", "Valor": "Arquivo não encontrado"})
            continue
        data = load_yaml(path)
        rows.extend(_achatar_config(nome, data, str(path)))
    return pd.DataFrame(rows, columns=["Arquivo", "Grupo", "Chave", "Valor"])


def _achatar_config(grupo: str, value: object, arquivo: str, prefixo: str = "") -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    if isinstance(value, dict):
        for key, item in value.items():
            chave = f"{prefixo}.{key}" if prefixo else str(key)
            rows.extend(_achatar_config(grupo, item, arquivo, chave))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            chave = f"{prefixo}[{index}]"
            rows.extend(_achatar_config(grupo, item, arquivo, chave))
    else:
        rows.append({"Arquivo": arquivo, "Grupo": grupo, "Chave": prefixo, "Valor": "" if value is None else str(value)})
    return rows


def gerar_excel(
    resultado_df: pd.DataFrame,
    municipios_df: pd.DataFrame,
    pendencias_df: pd.DataFrame,
    caminho_saida: str | Path | None = None,
    fontes_log_df: pd.DataFrame | None = None,
    configuracao_df: pd.DataFrame | None = None,
) -> Path:
    output_path = Path(caminho_saida) if caminho_saida else project_path("data", "output", "resultado_final_point_c.xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    resultado = garantir_colunas(resultado_df, COLUNAS_RESULTADO)
    resultado = normalizar_status_sem_dado_publicado(resultado)
    municipios = garantir_colunas(municipios_df, COLUNAS_MUNICIPIOS)
    pendencias = garantir_colunas(pendencias_df, COLUNAS_PENDENCIAS)
    fontes_log = garantir_colunas(fontes_log_df if fontes_log_df is not None else pd.DataFrame(), COLUNAS_FONTES_LOG)
    configuracao = configuracao_df if configuracao_df is not None else criar_configuracao_df()
    resumo = criar_resumo(resultado, municipios, pendencias)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        resultado.to_excel(writer, sheet_name="Dados", index=False)
        resultado.to_excel(writer, sheet_name="Resultado consolidado", index=False)
        municipios.to_excel(writer, sheet_name="Municípios pesquisados", index=False)
        pendencias.to_excel(writer, sheet_name="Pendências", index=False)
        resumo.to_excel(writer, sheet_name="Resumo", index=False)
        fontes_log.to_excel(writer, sheet_name="Fontes e Log", index=False)
        configuracao.to_excel(writer, sheet_name="Configuração", index=False)

        for sheet_name in writer.book.sheetnames:
            _formatar_aba(writer, sheet_name)

    return output_path
